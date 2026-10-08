#!/usr/bin/env python3
"""Genera la planilla parametrizada de liquidacion de comisiones (Fase 1).

Uso:
    python -I scripts/comisiones/build_plantilla_liquidacion.py --salida <out.xlsx>
    python -I scripts/comisiones/build_plantilla_liquidacion.py --salida <out.xlsx> --origen <planilla_vieja.xlsx>

Sin --origen genera la plantilla vacia (7 vendedores, una fila en blanco por tabla).
Con --origen importa las filas de la planilla vieja EN TIEMPO DE EJECUCION y valida
que el TOTAL A COBRAR recalculado coincida con el de la planilla original.

En este archivo solo hay PARAMETROS (porcentajes, gastos, nombres de vendedores).
Ningun dato de clientes queda versionado: las filas salen siempre del --origen.

El diseno de cada hoja (celdas fijas, encabezados de la tabla) es el de la seccion D
del spec compartido con el exportador TS de la Fase 2: no mover celdas sin cambiar
tambien el exportador.
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import re
import sys
import unicodedata
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

# --------------------------------------------------------------------------------------
# PARAMETROS (lo unico que vive en el repo)
# --------------------------------------------------------------------------------------
IVA_DIV = 1.10

GADM = [  # Movimiento (M) -> gasto administrativo
    ("V", 30000), ("RI", 30000), ("SAA", 30000), ("INC", 15000),
    ("CP", 0), ("CONT", 0), ("RETENCION", 0),
]
PCT_EXCEPCION = [  # Plan o M -> % que pisa el % del vendedor
    ("PM", 0.0), ("PO", 0.0), ("CP", 0.0),
    ("CONT", 0.20), ("SAA", 0.30), ("RETENCION", 0.20),
]
ADICIONAL_PLAN = [("PM", 350000)]  # Plan -> monto fijo por fila

# hoja = nombre de la hoja (igual al de la planilla original, para poder comparar)
# nombre = nombre del vendedor tal como figura en la celda de nombre de la hoja original
SALESPEOPLE = [
    {"hoja": "Rossana", "nombre": "ROSSANA PALACIOS", "pct": {"INDIVIDUAL": 1.00}},
    {"hoja": "Antonio P.", "nombre": "Antonio Paiva", "pct": {"INDIVIDUAL": 0.80}},
    {"hoja": "Catherine", "nombre": "Catherine", "pct": {"INDIVIDUAL": 0.50}},
    {"hoja": "SARA", "nombre": "Sara López", "pct": {"INDIVIDUAL": 0.50, "GRUPAL": 0.30}},
    {"hoja": "RODRIGO AMARILLA", "nombre": "RODRIGO AMARILLA", "pct": {"INDIVIDUAL": 0.50}},
    {"hoja": "Jorge Galeano", "nombre": "Jorge Galeano", "pct": {"INDIVIDUAL": 0.20}},
    {"hoja": "Olivia", "nombre": "OLIVIA OLMEDO", "pct": {"INDIVIDUAL": 0.50}},
]

# Capacidad de las tablas de Parametros (filas en blanco para que el usuario agregue)
PARAM_CAPACITY = 30
VEND_CAPACITY = 20

# --------------------------------------------------------------------------------------
# LAYOUT (spec seccion D) -- compartido con src/lib/commissions/exportLiquidacionXlsx.ts
# --------------------------------------------------------------------------------------
HEADER_ROW = 13
FIRST_DATA_ROW = 14
COLUMNS = [
    "Bloque", "Rec", "Fec", "Nombre", "Plan", "M", "Cto N°", "Vidas", "Total",
    "G Adm", "Cuota", "Cuota-IVA", "%", "Comisión", "Adicional", "Obs",
    "G Adm manual", "% manual", "Adicional manual",
]
COL = {name: get_column_letter(i + 1) for i, name in enumerate(COLUMNS)}
MONEY_COLS = ["Total", "G Adm", "Cuota", "Cuota-IVA", "Comisión", "Adicional",
              "G Adm manual", "Adicional manual"]
PCT_COLS = ["%", "% manual"]
MANUAL_COLS = ["G Adm manual", "% manual", "Adicional manual"]
INPUT_COLS = ["Bloque", "Rec", "Fec", "Nombre", "Plan", "M", "Cto N°", "Vidas", "Total", "Obs"]

FMT_GS = "#,##0"
FMT_PCT = "0.00%"
FMT_DATE = "dd/mm/yyyy"

FILL_INPUT = PatternFill("solid", fgColor="FFF9C4")    # amarillo claro: insumo
FILL_MANUAL = PatternFill("solid", fgColor="FFD8B0")   # naranja claro: pisa un parametro
FILL_HEAD = PatternFill("solid", fgColor="DDE7F3")
FILL_TOTAL = PatternFill("solid", fgColor="C8E6C9")
FILL_ALERT = PatternFill("solid", fgColor="FFB3B3")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="B0B0B0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


# --------------------------------------------------------------------------------------
# Utilidades
# --------------------------------------------------------------------------------------
def norm(v) -> str:
    """Mayusculas, sin acentos, espacios colapsados."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip().upper()


def table_key(sheet_name: str) -> str:
    s = unicodedata.normalize("NFKD", sheet_name)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^A-Za-z0-9]", "", s) or "Hoja"


def fmt_gs(n: float) -> str:
    return f"{n:,.0f}".replace(",", ".")


def fmt_pct(p: float) -> str:
    return f"{p * 100:g}%"


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    return None


def text(v):
    if v is None or isinstance(v, (int, float)):
        return None
    s = str(v).strip()
    return s or None


def sref(table: str, col: str) -> str:
    """Referencia estructurada a la fila actual, en formato de archivo."""
    esc = re.sub(r"(['\[\]#])", r"'\1", col)
    return f"{table}[[#This Row],[{esc}]]"


def cref(table: str, col: str) -> str:
    esc = re.sub(r"(['\[\]#])", r"'\1", col)
    return f"{table}[{esc}]"


def sheet_quote(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


# --------------------------------------------------------------------------------------
# Modelo de datos
# --------------------------------------------------------------------------------------
@dataclass
class Params:
    iva_div: float = IVA_DIV
    gadm: dict = field(default_factory=lambda: {k: v for k, v in GADM})
    pct_exc: dict = field(default_factory=lambda: {k: v for k, v in PCT_EXCEPCION})
    adic_plan: dict = field(default_factory=lambda: {k: v for k, v in ADICIONAL_PLAN})
    pct_vend: dict = field(default_factory=dict)  # (NOMBRE_NORM, BLOQUE) -> pct

    # Replica exacta de las formulas de la plantilla -----------------------------
    def rule_gadm(self, m):
        return self.gadm.get(norm(m), 0) if text(m) else 0

    def rule_pct(self, vendor, bloque, plan, m):
        if text(plan) and norm(plan) in self.pct_exc:
            return self.pct_exc[norm(plan)]
        if text(m) and norm(m) in self.pct_exc:
            return self.pct_exc[norm(m)]
        return self.pct_vend.get((norm(vendor), norm(bloque)), 0)

    def rule_adic(self, plan):
        return self.adic_plan.get(norm(plan), 0) if text(plan) else 0


@dataclass
class Row:
    bloque: str = "INDIVIDUAL"
    rec: object = None
    fec: object = None
    nombre: object = None
    plan: object = None
    m: object = None
    cto: object = None
    vidas: object = None
    total: object = None
    obs: list = field(default_factory=list)
    gadm_manual: object = None
    pct_manual: object = None
    adic_manual: object = None
    # valores de la planilla original (solo para validar)
    orig_comision: object = None
    orig_adicional: object = None


@dataclass
class Footer:
    viatico: float = 0
    recupero: float = 0
    bonif_base: float = 0
    bonif_pct: float = 0
    adicional_extra: float = 0
    adicional_extra_nota: str | None = None
    descuentos: float = 0
    # solo para validar
    orig_adicional: float | None = None
    orig_total: float | None = None
    orig_total_fallback: float | None = None


@dataclass
class Seller:
    hoja: str
    nombre: str
    pct: dict
    periodo: str = ""
    rows: list = field(default_factory=list)
    footer: Footer = field(default_factory=Footer)
    warnings: list = field(default_factory=list)

    @property
    def table(self):
        return "T_" + table_key(self.hoja)


def model_sheet(s: Seller, p: Params) -> dict:
    """Calcula en Python lo mismo que las formulas de la hoja."""
    com_total = vid = ventas = adic = ind = grp = 0.0
    per_row = []
    for r in s.rows:
        total = num(r.total) or 0
        gadm = r.gadm_manual if r.gadm_manual is not None else p.rule_gadm(r.m)
        cuota = total - gadm
        civa = cuota / p.iva_div
        pct = r.pct_manual if r.pct_manual is not None else p.rule_pct(s.nombre, r.bloque, r.plan, r.m)
        com = civa * pct
        ad = r.adic_manual if r.adic_manual is not None else p.rule_adic(r.plan)
        per_row.append((com, ad))
        com_total += com
        adic += ad
        vid += num(r.vidas) or 0
        ventas += total
        if norm(r.bloque) == "INDIVIDUAL":
            ind += com
        elif norm(r.bloque) == "GRUPAL":
            grp += com
    f = s.footer
    bonif = f.bonif_base * f.bonif_pct
    adic_total = adic + f.adicional_extra
    total = com_total + f.viatico + f.recupero + bonif + adic_total - f.descuentos
    return {"E4": vid, "E5": ventas, "E6": com_total, "E7": bonif, "E8": adic_total,
            "E9": ind, "E10": grp, "E11": total, "rows": per_row}


# --------------------------------------------------------------------------------------
# Importacion de la planilla vieja
# --------------------------------------------------------------------------------------
def classify_header(v):
    if v is None:
        return None
    raw = str(v).strip()
    if raw == "%":
        return "pct"
    k = re.sub(r"[^A-Z0-9]", "", norm(raw))
    if k.startswith("REC"):
        return "rec"
    if k.startswith("CTO"):
        return "cto"
    if k == "CUOTAIVA":
        return "cuotaiva"
    return {"FEC": "fec", "NOMBRE": "nombre", "PLAN": "plan", "M": "m", "VIDAS": "vidas",
            "TOTAL": "total", "GADM": "gadm", "CUOTA": "cuota", "BONIF": "bonif",
            "COMISION": "comision"}.get(k)


def import_seller(cfg: dict, ws, wv, params: Params) -> Seller:
    s = Seller(hoja=cfg["hoja"], nombre=cfg["nombre"], pct=cfg["pct"])
    max_c = min(ws.max_column, 40)

    # 1. fila de encabezado: la que tiene "Nombre"
    header = None
    for r in range(1, min(ws.max_row, 30) + 1):
        if any(norm(ws.cell(r, c).value) == "NOMBRE" for c in range(1, max_c + 1)):
            header = r
            break
    if header is None:
        raise SystemExit(f"[{s.hoja}] no encontre la fila de encabezado con 'Nombre'")
    cols = {}
    for c in range(1, max_c + 1):
        k = classify_header(ws.cell(header, c).value)
        if k and k not in cols:
            cols[k] = c
    for need in ("nombre", "total", "comision", "fec"):
        if need not in cols:
            raise SystemExit(f"[{s.hoja}] falta la columna '{need}' en el encabezado (fila {header})")
    cols.setdefault("rec", cols["fec"] - 1)          # Rossana no tiene titulo "Rec N°"
    c_adic = cols["comision"] + 1                    # adicional: columna a la derecha de Comision
    c_dif = c_adic + 1                               # etiqueta "DIF"

    # 2. nombre del vendedor y periodo (arriba del encabezado)
    period_date = None
    for r in range(1, header):
        for c in range(1, max_c + 1):
            vv = ws.cell(r, c).value
            if isinstance(vv, dt.datetime) and period_date is None:
                period_date = vv
    name_cell = None
    for r in range(1, header):
        v = ws.cell(r, cols["nombre"]).value
        if isinstance(v, str) and v.strip():
            name_cell = v.strip()
            break
    if name_cell and name_cell != cfg["nombre"]:
        s.warnings.append(f"nombre en la planilla '{name_cell}' != parametro '{cfg['nombre']}'; se usa el de la planilla")
        s.nombre = name_cell
        params.pct_vend.update({(norm(name_cell), b): v for b, v in cfg["pct"].items()})
    if period_date:
        last = calendar.monthrange(period_date.year, period_date.month)[1]
        s.periodo = f"01/{period_date.month:02d}/{period_date.year} – {last:02d}/{period_date.month:02d}/{period_date.year}"
        period_month = MESES[period_date.month - 1]
    else:
        period_month = None

    def cv(r, c):
        return wv.cell(r, c).value

    def row_is_header(r):
        return norm(ws.cell(r, cols["nombre"]).value) == "NOMBRE"

    def row_has_sum(r):
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.lstrip("=+").upper().startswith("SUM("):
                return True
        return False

    # 3. filas de datos
    bloque = "INDIVIDUAL"
    in_section = True
    first_end = None
    r = header + 1
    while r <= ws.max_row:
        if in_section:
            name = text(ws.cell(r, cols["nombre"]).value)
            if row_is_header(r):
                r += 1
                continue
            if name is None or row_has_sum(r):
                in_section = False
                first_end = first_end or r
                r += 1
                continue
            row = Row(bloque=bloque)
            row.rec = text(ws.cell(r, cols["rec"]).value)
            fec = ws.cell(r, cols["fec"]).value
            row.fec = fec.date() if isinstance(fec, dt.datetime) else text(fec)
            row.nombre = name
            row.plan = text(cv(r, cols["plan"])) if "plan" in cols else None
            row.m = text(cv(r, cols["m"])) if "m" in cols else None
            row.cto = num(cv(r, cols["cto"])) if "cto" in cols else None
            row.vidas = num(cv(r, cols["vidas"])) if "vidas" in cols else None
            row.total = num(cv(r, cols["total"]))
            gadm_o = (num(cv(r, cols["gadm"])) if "gadm" in cols else None) or 0
            pct_o = (num(cv(r, cols["pct"])) if "pct" in cols else None) or 0
            adic_v = cv(r, c_adic)
            adic_o = num(adic_v) or 0
            if text(adic_v):
                row.obs.append(text(adic_v))
            if text(cv(r, c_dif)):
                row.obs.append(text(cv(r, c_dif)))
            row.orig_comision = num(cv(r, cols["comision"])) or 0
            row.orig_adicional = adic_o

            # overrides: solo si el original difiere de lo que daria la regla
            total = row.total or 0
            if abs(gadm_o - params.rule_gadm(row.m)) > 0.005:
                row.gadm_manual = gadm_o
            rule_pct = params.rule_pct(s.nombre, bloque, row.plan, row.m)
            if abs(pct_o - rule_pct) > 1e-9:
                row.pct_manual = pct_o
                if 0 < pct_o < 0.05:
                    row.obs.append(f"VERIFICAR: {fmt_pct(pct_o)} sobre {fmt_gs(total)}")
                else:
                    row.obs.append(f"VERIFICAR: % {fmt_pct(pct_o)} distinto de la regla ({fmt_pct(rule_pct)})")
            if abs(adic_o - params.rule_adic(row.plan)) > 0.005:
                row.adic_manual = adic_o

            # avisos
            if isinstance(row.fec, str) and period_month and norm(row.fec) in MESES and norm(row.fec) != period_month:
                row.obs.append(f"VERIFICAR: Fec {row.fec} en liquidación de {period_month.lower()}")
            if "cuota" in cols:
                cuota_o = num(cv(r, cols["cuota"]))
                if cuota_o is not None and abs(cuota_o - (total - gadm_o)) > 0.5:
                    row.obs.append(f"Original: Cuota cargada a mano en {fmt_gs(cuota_o)}")
                    s.warnings.append(f"fila {r} ({name}): Cuota original {fmt_gs(cuota_o)} != Total-G Adm {fmt_gs(total - gadm_o)}")
            s.rows.append(row)
        else:
            if any(norm(ws.cell(r, c).value) == "EMPRESARIALES" for c in range(1, max_c + 1)):
                bloque = "GRUPAL"
                in_section = True
        r += 1

    # 4. pie: etiquetas a la izquierda de la columna Comision, valor en la columna Comision
    f = s.footer
    c_com = cols["comision"]
    start = first_end or (header + 1)
    for r in range(start, ws.max_row + 1):
        for c in range(max(1, c_com - 4), c_com):
            raw = ws.cell(r, c).value
            if not isinstance(raw, str) or raw.lstrip().startswith("="):
                continue  # solo etiquetas de texto
            lab = norm(raw)
            if not lab:
                continue
            val = cv(r, c_com)
            v = num(val) or 0
            if lab == "VIATICO":
                f.viatico += v
            elif lab in ("RECUPERO", "RECUPERO CUOTAS"):
                f.recupero += v
            elif lab.startswith("BONIFICACION"):
                formula = ws.cell(r, c_com).value
                m = re.match(r"^=\s*\+?\s*\$?([A-Z]{1,3})\$?(\d+)\s*\*\s*([\d.,]+)\s*%\s*$", str(formula or ""))
                if m:
                    f.bonif_base = num(wv[f"{m.group(1)}{m.group(2)}"].value) or 0
                    f.bonif_pct = float(m.group(3).replace(",", ".")) / 100
                else:
                    mp = re.search(r"(\d+(?:[.,]\d+)?)\s*%", lab)
                    pct = float(mp.group(1).replace(",", ".")) / 100 if mp else 1
                    f.bonif_base, f.bonif_pct = (v / pct if pct else v), pct
            elif lab in ("AJUSTE VALE", "OTROS"):
                if v:
                    f.adicional_extra += v
                    nota = f"Importado de '{ws.cell(r, c).value.strip()}' ({fmt_gs(v)}) de la planilla original."
                    f.adicional_extra_nota = ((f.adicional_extra_nota + " ") if f.adicional_extra_nota else "") + nota
            elif lab == "DESCUENTOS":
                f.descuentos += v
            elif lab == "ADICIONAL":
                f.orig_adicional = (f.orig_adicional or 0) + v
            elif lab == "TOTAL A COBRAR":
                f.orig_total = num(val)
            elif lab == "TOTAL" and f.orig_total is None:
                f.orig_total_fallback = num(val)
            elif lab in ("TOTALES", "COMISION"):
                pass
            elif num(val):
                s.warnings.append(f"pie: etiqueta '{ws.cell(r, c).value}' con valor {val} no reconocida (ignorada)")
    if f.orig_total is None:
        f.orig_total = f.orig_total_fallback
    if f.orig_adicional is not None:
        sum_rows = sum(r.orig_adicional or 0 for r in s.rows)
        if abs(sum_rows - f.orig_adicional) > 0.5:
            s.warnings.append(f"ADICIONAL del pie ({fmt_gs(f.orig_adicional)}) != suma de adicionales por fila ({fmt_gs(sum_rows)})")
    return s


# --------------------------------------------------------------------------------------
# Escritura del libro
# --------------------------------------------------------------------------------------
def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name=name, attr_text=ref)


def build_parametros(wb, sellers, params: Params):
    ws = wb.create_sheet("Parametros")
    ws["A1"] = "PARÁMETROS DE LA LIQUIDACIÓN"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Divisor IVA"
    ws["A3"].font = BOLD
    ws["B3"] = params.iva_div
    ws["B3"].fill = FILL_INPUT
    ws["B3"].border = BOX
    ws["B3"].number_format = "0.00"
    ws["C3"] = "IVA_DIV. Cuota-IVA = Cuota / IVA_DIV. 1,10 = IVA del 10% incluido en la cuota."
    add_name(wb, "IVA_DIV", "Parametros!$B$3")

    first = 7
    last = first + PARAM_CAPACITY - 1
    tables = [
        # (col inicial, nombre, encabezados, filas, formatos, nota)
        ("A", "tbl_GAdm", ["M", "G Adm"], [list(x) for x in params.gadm.items()], [None, FMT_GS],
         "Gasto administrativo por Movimiento (columna M). Un M que no está en la lista da 0."),
        ("D", "tbl_PctExcepcion", ["Clave", "%"], [list(x) for x in params.pct_exc.items()], [None, FMT_PCT],
         "Clave = Plan o M. Pisa el % del vendedor. Se busca primero el Plan y después el M."),
        ("G", "tbl_AdicionalPlan", ["Plan", "Adicional"], [list(x) for x in params.adic_plan.items()], [None, FMT_GS],
         "Monto fijo que se suma por fila según el Plan (ej. prima de Plan Materno)."),
        ("J", "tbl_PctVendedor", ["Vendedor", "Bloque", "%"],
         [[s.nombre, b, pct] for s in sellers for b, pct in s.pct.items()], [None, None, FMT_PCT],
         "% del vendedor por Bloque (INDIVIDUAL / GRUPAL). Vendedor = celda B2 de su hoja."),
        ("N", "tbl_Vendedores", ["Vendedor", "Hoja"], [[s.nombre, s.hoja] for s in sellers], [None, None],
         "Una fila por vendedor. Hoja = nombre exacto de la pestaña. Alimenta el Resumen."),
    ]
    dv_bloque = DataValidation(type="list", formula1='"INDIVIDUAL,GRUPAL"', allow_blank=True)
    ws.add_data_validation(dv_bloque)
    for start, name, heads, rows, fmts, note in tables:
        c0 = ws[f"{start}1"].column
        cap_last = first + (VEND_CAPACITY if name == "tbl_Vendedores" else PARAM_CAPACITY) - 1
        if len(rows) > cap_last - first + 1:
            raise SystemExit(f"{name}: más filas que la capacidad")
        ws.cell(4, c0, name).font = BOLD
        ws.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c0 + len(heads) - 1)
        ws.cell(5, c0, note).alignment = WRAP
        ws.cell(5, c0).font = Font(italic=True, size=9, color="555555")
        for j, h in enumerate(heads):
            cell = ws.cell(6, c0 + j, h)
            cell.font = BOLD
            cell.fill = FILL_HEAD
            cell.border = BOX
        for i in range(first, cap_last + 1):
            for j in range(len(heads)):
                cell = ws.cell(i, c0 + j)
                cell.fill = FILL_INPUT
                cell.border = BOX
                if fmts[j]:
                    cell.number_format = fmts[j]
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.cell(first + i, c0 + j, v)
        cl = [get_column_letter(c0 + j) for j in range(len(heads))]
        add_name(wb, name, f"Parametros!${cl[0]}${first}:${cl[-1]}${cap_last}")
        if name == "tbl_PctExcepcion":
            add_name(wb, "PctExc_Clave", f"Parametros!${cl[0]}${first}:${cl[0]}${cap_last}")
        if name == "tbl_PctVendedor":
            add_name(wb, "PctV_Vendedor", f"Parametros!${cl[0]}${first}:${cl[0]}${cap_last}")
            add_name(wb, "PctV_Bloque", f"Parametros!${cl[1]}${first}:${cl[1]}${cap_last}")
            add_name(wb, "PctV_Pct", f"Parametros!${cl[2]}${first}:${cl[2]}${cap_last}")
            dv_bloque.add(f"{cl[1]}{first}:{cl[1]}{cap_last}")
        if name == "tbl_Vendedores":
            add_name(wb, "Vend_Nombre", f"Parametros!${cl[0]}${first}:${cl[0]}${cap_last}")
            add_name(wb, "Vend_Hoja", f"Parametros!${cl[1]}${first}:${cl[1]}${cap_last}")
    ws.row_dimensions[5].height = 60
    for col, w in {"A": 14, "B": 12, "C": 4, "D": 14, "E": 10, "F": 4, "G": 12, "H": 12, "I": 4,
                   "J": 24, "K": 13, "L": 9, "M": 4, "N": 24, "O": 22}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A7"
    return ws


def row_formulas(t: str) -> dict:
    """Formulas por fila en formato de archivo (referencias estructuradas)."""
    R = lambda c: sref(t, c)  # noqa: E731
    return {
        "G Adm": f'IF({R("G Adm manual")}<>"",{R("G Adm manual")},IFERROR(VLOOKUP({R("M")},tbl_GAdm,2,FALSE),0))',
        "Cuota": f'{R("Total")}-{R("G Adm")}',
        "Cuota-IVA": f'{R("Cuota")}/IVA_DIV',
        "%": (f'IF({R("% manual")}<>"",{R("% manual")},'
              f'IFERROR(VLOOKUP({R("Plan")},tbl_PctExcepcion,2,FALSE),'
              f'IFERROR(VLOOKUP({R("M")},tbl_PctExcepcion,2,FALSE),'
              f'SUMIFS(PctV_Pct,PctV_Vendedor,$B$2,PctV_Bloque,{R("Bloque")}))))'),
        "Comisión": f'{R("Cuota-IVA")}*{R("%")}',
        "Adicional": f'IF({R("Adicional manual")}<>"",{R("Adicional manual")},IFERROR(VLOOKUP({R("Plan")},tbl_AdicionalPlan,2,FALSE),0))',
    }


def build_seller_sheet(wb, s: Seller):
    ws = wb.create_sheet(s.hoja)
    t = s.table
    ws["A1"] = "LIQUIDACIÓN DE COMISIONES"
    ws["A1"].font = Font(bold=True, size=14)
    labels_a = {2: "Vendedor", 3: "Período", 4: "Viático", 5: "Recupero", 6: "Base bonificación",
                7: "% bonificación", 8: "Adicional extra", 9: "Descuentos (positivo)"}
    for r, lab in labels_a.items():
        ws.cell(r, 1, lab).font = BOLD
    ws["B2"] = s.nombre
    ws["B2"].comment = Comment("Debe coincidir con 'Vendedor' en tbl_PctVendedor y tbl_Vendedores (hoja Parametros).", "Plantilla")
    ws["B3"] = s.periodo
    ws["D3"] = "Liquidación"
    ws["D3"].font = BOLD
    f = s.footer
    inputs = {"B4": f.viatico, "B5": f.recupero, "B6": f.bonif_base, "B7": f.bonif_pct,
              "B8": f.adicional_extra, "B9": f.descuentos}
    for ref, v in inputs.items():
        ws[ref] = v
        ws[ref].number_format = FMT_PCT if ref == "B7" else FMT_GS
    for ref in ["B2", "B3", "E3", *inputs]:
        ws[ref].fill = FILL_INPUT
        ws[ref].border = BOX
    if f.adicional_extra_nota:
        ws["B8"].comment = Comment(f.adicional_extra_nota, "Importación")

    calc = {
        4: ("Vidas", f"=SUM({cref(t, 'Vidas')})", "#,##0"),
        5: ("Ventas (Total)", f"=SUM({cref(t, 'Total')})", FMT_GS),
        6: ("Comisión", f"=SUM({cref(t, 'Comisión')})", FMT_GS),
        7: ("Bonificación", "=B6*B7", FMT_GS),
        8: ("Adicional total", f"=SUM({cref(t, 'Adicional')})+B8", FMT_GS),
        9: ("Subtotal Individual", f'=SUMIF({cref(t, "Bloque")},"INDIVIDUAL",{cref(t, "Comisión")})', FMT_GS),
        10: ("Subtotal Grupal", f'=SUMIF({cref(t, "Bloque")},"GRUPAL",{cref(t, "Comisión")})', FMT_GS),
    }
    for r, (lab, formula, fmt) in calc.items():
        ws.cell(r, 4, lab).font = BOLD
        c = ws.cell(r, 5, formula)
        c.number_format = fmt
        c.border = BOX
    ws["A11"] = "TOTAL A COBRAR"
    ws["E11"] = "=E6+B4+B5+E7+E8-B9"
    for ref in ("A11", "E11"):
        ws[ref].font = Font(bold=True, size=12)
        ws[ref].fill = FILL_TOTAL
    ws["E11"].number_format = FMT_GS
    ws["E11"].border = BOX

    # Tabla
    formulas = row_formulas(t)
    for j, h in enumerate(COLUMNS):
        cell = ws.cell(HEADER_ROW, j + 1, h)
        cell.font = BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    rows = s.rows or [Row(bloque=None)]
    for i, row in enumerate(rows):
        r = FIRST_DATA_ROW + i
        values = {
            "Bloque": row.bloque, "Rec": row.rec, "Fec": row.fec, "Nombre": row.nombre,
            "Plan": row.plan, "M": row.m, "Cto N°": row.cto, "Vidas": row.vidas,
            "Total": row.total, "Obs": " | ".join(row.obs) if row.obs else None,
            "G Adm manual": row.gadm_manual, "% manual": row.pct_manual,
            "Adicional manual": row.adic_manual,
        }
        for j, h in enumerate(COLUMNS):
            cell = ws.cell(r, j + 1)
            if h in formulas:
                cell.value = "=" + formulas[h]
            else:
                cell.value = values[h]
            if h in MONEY_COLS:
                cell.number_format = FMT_GS
            elif h in PCT_COLS:
                cell.number_format = FMT_PCT
            elif h == "Fec":
                cell.number_format = FMT_DATE
            elif h in ("Cto N°", "Vidas"):
                cell.number_format = "0"
            if h in MANUAL_COLS:
                cell.fill = FILL_MANUAL
            if h == "Obs":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = FIRST_DATA_ROW + len(rows) - 1
    ref = f"A{HEADER_ROW}:{COL['Adicional manual']}{last}"
    tcols = []
    for j, h in enumerate(COLUMNS):
        tc = TableColumn(id=j + 1, name=h)
        if h in formulas:
            tc.calculatedColumnFormula = TableFormula(attr_text=formulas[h])
        tcols.append(tc)
    tab = Table(displayName=t, name=t, ref=ref, tableColumns=tcols, autoFilter=AutoFilter(ref=ref),
                tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True))
    ws.add_table(tab)
    for h in MANUAL_COLS:
        ws[f"{COL[h]}{HEADER_ROW}"].fill = FILL_MANUAL
        ws[f"{COL[h]}{HEADER_ROW}"].comment = Comment(
            "Columna manual: si tiene un valor, pisa el parámetro para esa fila. Vacía = regla de Parametros.", "Plantilla")

    # validacion de Bloque y alerta de % sin regla
    dv = DataValidation(type="list", formula1='"INDIVIDUAL,GRUPAL"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A{FIRST_DATA_ROW}:A1000")
    d, m_, e, fcol, r_, a = COL["Nombre"], COL["%"], COL["Plan"], COL["M"], COL["% manual"], COL["Bloque"]
    rule = FormulaRule(
        formula=[f'AND(${d}{FIRST_DATA_ROW}<>"",${r_}{FIRST_DATA_ROW}="",'
                 f'ISNA(MATCH(${e}{FIRST_DATA_ROW},PctExc_Clave,0)),ISNA(MATCH(${fcol}{FIRST_DATA_ROW},PctExc_Clave,0)),'
                 f'COUNTIFS(PctV_Vendedor,$B$2,PctV_Bloque,${a}{FIRST_DATA_ROW})=0)'],
        fill=FILL_ALERT)
    ws.conditional_formatting.add(f"{m_}{FIRST_DATA_ROW}:{m_}1000", rule)

    widths = {"Bloque": 12, "Rec": 9, "Fec": 11, "Nombre": 28, "Plan": 10, "M": 6, "Cto N°": 9,
              "Vidas": 6, "Total": 12, "G Adm": 10, "Cuota": 12, "Cuota-IVA": 12, "%": 8,
              "Comisión": 12, "Adicional": 11, "Obs": 34, "G Adm manual": 11, "% manual": 9,
              "Adicional manual": 11}
    for h, w in widths.items():
        ws.column_dimensions[COL[h]].width = w
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 14
    ws.row_dimensions[HEADER_ROW].height = 30
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    return ws


def build_resumen(wb, sellers):
    ws = wb.create_sheet("Resumen")
    ws["A1"] = "RESUMEN DE LIQUIDACIÓN DE COMISIONES"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = next((s.periodo for s in sellers if s.periodo), "")
    heads = ["Vendedor", "Hoja", "Vidas", "Ventas", "Comisión", "Bonificación", "Viático",
             "Recupero", "Adicional", "Descuentos", "Total", "Control"]
    for j, h in enumerate(heads):
        c = ws.cell(3, j + 1, h)
        c.font = BOLD
        c.fill = FILL_HEAD
        c.border = BOX
    src = {"C": "$E$4", "D": "$E$5", "E": "$E$6", "F": "$E$7", "G": "$B$4", "H": "$B$5",
           "I": "$E$8", "J": "$B$9"}
    first = 4
    last = first + VEND_CAPACITY - 1
    for k in range(VEND_CAPACITY):
        r = first + k
        p = 7 + k  # fila en tbl_Vendedores
        ws[f"A{r}"] = f'=IF(Parametros!$O${p}="","",Parametros!$N${p})'
        ws[f"B{r}"] = f'=IF(Parametros!$O${p}="","",Parametros!$O${p})'
        ind = lambda cell: f'INDIRECT("\'"&SUBSTITUTE($B{r},"\'","\'\'")&"\'!{cell}")'  # noqa: E731
        for col, cell in src.items():
            ws[f"{col}{r}"] = f'=IF($B{r}="","",{ind(cell)})'
        ws[f"K{r}"] = f'=IF($B{r}="","",E{r}+F{r}+G{r}+H{r}+I{r}-J{r})'
        ws[f"L{r}"] = f'=IF($B{r}="","",K{r}-{ind("$E$11")})'
        for col in "CDEFGHIJKL":
            ws[f"{col}{r}"].number_format = FMT_GS
    tr = last + 1
    ws[f"A{tr}"] = "TOTAL"
    for col in "CDEFGHIJKL":
        ws[f"{col}{tr}"] = f"=SUM({col}{first}:{col}{last})"
        ws[f"{col}{tr}"].number_format = FMT_GS
    for col in "ABCDEFGHIJKL":
        ws[f"{col}{tr}"].font = BOLD
        ws[f"{col}{tr}"].fill = FILL_TOTAL
    ws["N3"] = "Control = Total del Resumen − TOTAL A COBRAR de la hoja (E11). Tiene que dar 0."
    ws["N3"].font = Font(italic=True, size=9, color="555555")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20
    for col in "CDEFGHIJKL":
        ws.column_dimensions[col].width = 13
    ws.freeze_panes = "A4"
    return ws


LEEME = [
    ("PLANILLA DE LIQUIDACIÓN DE COMISIONES — CÓMO SE USA", "title"),
    ("", None),
    ("Colores", "h"),
    ("Amarillo claro = celda de insumo (se completa a mano). Naranja claro = columna manual que pisa un parámetro para esa fila. El resto son fórmulas: no se tocan.", None),
    ("", None),
    ("Cambiar parámetros (hoja Parametros)", "h"),
    ("• IVA: cambiar la celda B3 (nombre IVA_DIV). 1,10 = IVA 10% incluido. Recalcula todas las hojas.", None),
    ("• Gasto administrativo: tabla tbl_GAdm (columnas A:B). Una fila por Movimiento (M). Si aparece un M nuevo, escribirlo en la primera fila vacía. Un M que no figura da 0.", None),
    ("• % por vendedor: tabla tbl_PctVendedor (J:L). Vendedor (igual a la celda B2 de su hoja), Bloque (INDIVIDUAL o GRUPAL) y %. Un vendedor con dos bloques lleva dos filas (ej. Sara).", None),
    ("• Excepciones: tabla tbl_PctExcepcion (D:E). La Clave puede ser un Plan o un M; si coincide, pisa el % del vendedor. Primero se busca el Plan y después el M. Ej.: PM = 0%, CONT = 20%.", None),
    ("• Adicional por plan: tabla tbl_AdicionalPlan (G:H). Monto fijo que se suma en la columna Adicional de cada fila con ese Plan (ej. PM = 350.000).", None),
    ("• Las tablas tienen filas vacías al final: agregar ahí, sin insertar filas en el medio, y las fórmulas las toman solas.", None),
    ("", None),
    ("Pisar un valor en una fila puntual", "h"),
    ("• Columnas naranjas de la derecha de cada tabla: 'G Adm manual', '% manual', 'Adicional manual'. Si tienen un valor, se usa ese en lugar del parámetro. Vacías = regla de Parametros.", None),
    ("• Usarlas solo para casos que no siguen la regla y dejar el motivo en Obs. Así no se rompe ninguna fórmula.", None),
    ("• El % se pinta de rojo cuando la fila no tiene regla aplicable (ni excepción ni % del vendedor para ese Bloque): revisar Parametros o cargar el % manual.", None),
    ("", None),
    ("Agregar una fila de venta", "h"),
    ("• Escribir en la primera fila vacía debajo de la tabla (empezando por la columna A o D). La tabla se agranda sola y copia las fórmulas; los totales de arriba (E4:E11) la incluyen sin tocar nada.", None),
    ("• También se puede insertar una fila dentro de la tabla (clic derecho > Insertar > Filas de la tabla arriba).", None),
    ("• Bloque: INDIVIDUAL o GRUPAL (lista desplegable). Define qué % del vendedor se aplica.", None),
    ("", None),
    ("Pie de la hoja (bloque de arriba)", "h"),
    ("• B4 Viático, B5 Recupero, B6 Base de bonificación y B7 % de bonificación (Bonificación = B6 × B7), B8 Adicional extra (fuera de las filas), B9 Descuentos en positivo.", None),
    ("• TOTAL A COBRAR (E11) = Comisión + Viático + Recupero + Bonificación + Adicional total − Descuentos.", None),
    ("• Las celdas del bloque de arriba no se mueven: el Resumen las lee por dirección fija.", None),
    ("", None),
    ("Agregar un vendedor", "h"),
    ("1. Clic derecho en la pestaña de un vendedor > Mover o copiar > Crear una copia. Renombrar la pestaña.", None),
    ("2. En la copia: cambiar B2 (nombre del vendedor), borrar las filas de la tabla y los insumos B4:B9.", None),
    ("3. Renombrar la tabla: seleccionar una celda de la tabla > Diseño de tabla > Nombre de la tabla (ej. T_Juan). Las fórmulas se actualizan solas.", None),
    ("4. En Parametros: agregar una fila en tbl_Vendedores (Vendedor + nombre exacto de la pestaña) y una fila por Bloque en tbl_PctVendedor.", None),
    ("5. El Resumen toma la fila nueva automáticamente (hasta 20 vendedores). La columna Control tiene que dar 0.", None),
    ("", None),
    ("Resumen", "h"),
    ("• Una fila por vendedor de tbl_Vendedores. Cada valor se lee con INDIRECTO de la hoja indicada en la columna Hoja, así el nombre y la hoja no se desalinean.", None),
    ("• Control = Total del Resumen − TOTAL A COBRAR de la hoja. Distinto de 0 = hay que revisar.", None),
    ("", None),
    ("Esta planilla se genera con scripts/comisiones/build_plantilla_liquidacion.py (repo prepaga-digital).", "small"),
]


def build_leeme(wb):
    ws = wb.create_sheet("Leeme", 0)
    ws.column_dimensions["A"].width = 130
    for i, (line, kind) in enumerate(LEEME, start=1):
        c = ws.cell(i, 1, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "title":
            c.font = Font(bold=True, size=14)
        elif kind == "h":
            c.font = Font(bold=True, size=12, color="1F4E79")
        elif kind == "small":
            c.font = Font(italic=True, size=9, color="777777")
    return ws


def build_workbook(sellers, params: Params, path: str):
    wb = Workbook()
    wb.remove(wb.active)
    build_leeme(wb)
    build_parametros(wb, sellers, params)
    for s in sellers:
        build_seller_sheet(wb, s)
    build_resumen(wb, sellers)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


# --------------------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--salida", required=True, help="archivo .xlsx a generar")
    ap.add_argument("--origen", help="planilla vieja de la que importar las filas")
    args = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

    params = Params()
    sellers = [Seller(hoja=c["hoja"], nombre=c["nombre"], pct=dict(c["pct"])) for c in SALESPEOPLE]

    if args.origen:
        wb_f = load_workbook(args.origen)                    # formulas
        wb_v = load_workbook(args.origen, data_only=True)    # valores cacheados
        imp = Params(iva_div=params.iva_div)
        imp.pct_vend = {(norm(c["nombre"]), b): v for c in SALESPEOPLE for b, v in c["pct"].items()}
        imported = []
        for cfg in SALESPEOPLE:
            if cfg["hoja"] not in wb_f.sheetnames:
                raise SystemExit(f"la planilla de origen no tiene la hoja '{cfg['hoja']}'")
            imported.append(import_seller(cfg, wb_f[cfg["hoja"]], wb_v[cfg["hoja"]], imp))
        sellers = imported
    params.pct_vend = {(norm(s.nombre), b): v for s in sellers for b, v in s.pct.items()}

    build_workbook(sellers, params, args.salida)
    print(f"OK: {args.salida} ({len(sellers)} vendedores, {sum(len(s.rows) for s in sellers)} filas importadas)")

    if args.origen:
        ok = validate(sellers, params)
        return 0 if ok else 1
    return 0


def validate(sellers, params: Params) -> bool:
    print()
    print("VALIDACIÓN: TOTAL A COBRAR recalculado (fórmulas de la plantilla) vs. planilla original")
    print(f"{'Hoja':<18}{'Filas':>6}{'Esperado':>18}{'Original':>18}{'Dif':>12}  Estado")
    all_ok = True
    for s in sellers:
        mdl = model_sheet(s, params)
        orig = s.footer.orig_total
        diff = None if orig is None else mdl["E11"] - orig
        ok = diff is not None and abs(diff) < 1
        all_ok &= ok
        print(f"{s.hoja:<18}{len(s.rows):>6}{mdl['E11']:>18,.2f}{(orig if orig is not None else float('nan')):>18,.2f}"
              f"{(diff if diff is not None else float('nan')):>12,.4f}  {'OK' if ok else 'DIFERENCIA'}")
        # chequeo fila por fila contra la comision/adicional cacheados del original
        for r, (com, ad) in zip(s.rows, mdl["rows"]):
            if abs(com - (r.orig_comision or 0)) > 0.01 or abs(ad - (r.orig_adicional or 0)) > 0.01:
                print(f"    fila '{r.nombre}': comisión {com:,.2f} vs {r.orig_comision:,.2f}; adicional {ad:,.0f} vs {r.orig_adicional:,.0f}")
                all_ok = False
        for w in s.warnings:
            print(f"    aviso: {w}")
    print("RESULTADO:", "todas las hojas cuadran" if all_ok else "HAY DIFERENCIAS")

    # sensibilidad: el IVA es un parametro de verdad
    p2 = Params(iva_div=1.0, pct_vend=params.pct_vend)
    print()
    print("Prueba de parámetro: IVA_DIV 1,10 -> 1,00")
    for s in sellers:
        a, b = model_sheet(s, params)["E11"], model_sheet(s, p2)["E11"]
        print(f"  {s.hoja:<18}{a:>16,.2f} -> {b:>16,.2f}")
    return all_ok


if __name__ == "__main__":
    sys.exit(main())
